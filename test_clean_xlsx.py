"""Run with python test_clean_xlsx.py from this folder."""

import csv
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from clean_xlsx import clean_xlsx


def book(path, rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)


with tempfile.TemporaryDirectory() as folder:
    folder = Path(folder)
    source, reference = folder / "source.xlsx", folder / "reference.xlsx"
    output, report = folder / "cleaned.xlsx", folder / "cleaned_report.csv"
    book(source, [["id", "business", "email"],
                  ["001", " Acme ", ""], ["002", "Blue", "old@example.com"]])
    book(reference, [["id", "business", "email", "segment"],
                     ["001", "Acme", "a@example.com", "Home"],
                     ["002", "Blue", "new@example.com", "Trade"]])
    before = source.read_bytes()
    stats = clean_xlsx(source, output, key_columns=["id"], reference=reference,
                       add_columns=["segment"])
    rows = list(load_workbook(output, read_only=True).active.values)
    assert rows == [("id", "business", "email", "segment"),
                    ("001", "Acme", "a@example.com", "Home"),
                    ("002", "Blue", "old@example.com", "Trade")], rows
    with report.open(encoding="utf-8-sig", newline="") as handle:
        audit = list(csv.DictReader(handle))
    assert "conflicts" in audit[1]["review_note"]
    assert stats["kept_records_needing_review"] == 2
    assert source.read_bytes() == before
    try:
        clean_xlsx(source, output)
        raise AssertionError("Existing output overwritten")
    except FileExistsError:
        pass

    formula = folder / "formula.xlsx"
    book(formula, [["id", "notes"], ["1", "=1+1"]])
    try:
        clean_xlsx(formula, folder / "bad.xlsx")
        raise AssertionError("Formula accepted")
    except ValueError:
        assert not (folder / "bad.xlsx").exists()

    formatted = folder / "formatted.xlsx"
    book(formatted, [["id"], [123]])
    workbook = load_workbook(formatted)
    workbook.active["A2"].number_format = "00000"
    workbook.save(formatted)
    try:
        clean_xlsx(formatted, folder / "formatted_out.xlsx")
        raise AssertionError("Formatted number silently changed")
    except ValueError:
        assert not (folder / "formatted_out.xlsx").exists()

    multi = folder / "multi.xlsx"
    book(multi, [["id"], ["1"]])
    workbook = load_workbook(multi)
    workbook.create_sheet("Other")
    workbook.save(multi)
    try:
        clean_xlsx(multi, folder / "multi_out.xlsx")
        raise AssertionError("Multiple sheets silently lost")
    except ValueError:
        assert not (folder / "multi_out.xlsx").exists()

print("PASS: Excel cleanup, exact-key enrichment, conflicts, formula and formatting gates, overwrite protection")
