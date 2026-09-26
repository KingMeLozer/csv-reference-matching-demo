"""Clean a small, data-only Excel workbook using the reviewed CSV workflow."""

import argparse
import csv
import io
import math
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from clean_csv import clean_csv


def _read_sheet(path):
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Input must be an .xlsx workbook")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Workbook must contain exactly one worksheet")
        sheet = workbook.active
        if sheet.max_row > 101 or sheet.max_column > 10:
            raise ValueError("Workbook exceeds the 100-row, 10-column demo limit")
        rows = []
        for row in sheet.iter_rows():
            values = []
            for cell in row:
                value = cell.value
                if value is None:
                    values.append("")
                elif cell.data_type in ("f", "e"):
                    raise ValueError(f"Formula or error cell at {sheet.title}!{cell.coordinate}")
                elif isinstance(value, str):
                    values.append(value)
                elif type(value) in (int, float) and cell.number_format == "General" and math.isfinite(value):
                    values.append(str(value))
                else:
                    raise ValueError(f"Unsupported formatted or typed cell at {sheet.title}!{cell.coordinate}")
            rows.append(values)
        return rows
    finally:
        workbook.close()


def _csv_bytes(rows):
    buffer = io.StringIO(newline="")
    csv.writer(buffer).writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def clean_xlsx(source, destination, report=None, key_columns=None, reference=None, add_columns=None):
    source, destination = Path(source), Path(destination)
    reference = Path(reference) if reference else None
    report = Path(report) if report else destination.with_name(destination.stem + "_report.csv")
    if destination.suffix.lower() != ".xlsx" or report.suffix.lower() != ".csv":
        raise ValueError("Output must be .xlsx and report must be .csv")
    paths = [source, destination, report] + ([reference] if reference else [])
    if len({path.resolve() for path in paths}) != len(paths):
        raise ValueError("Source, reference, output, and report must be different files")
    if destination.exists() or report.exists():
        raise FileExistsError("Output or report exists; choose new filenames")

    source_rows = _read_sheet(source)
    reference_rows = _read_sheet(reference) if reference else None
    with tempfile.TemporaryDirectory() as folder:
        folder = Path(folder)
        csv_source, csv_output = folder / "source.csv", folder / "cleaned.csv"
        csv_source.write_bytes(_csv_bytes(source_rows))
        csv_reference = None
        if reference_rows is not None:
            csv_reference = folder / "reference.csv"
            csv_reference.write_bytes(_csv_bytes(reference_rows))
        stats = clean_csv(csv_source, csv_output, key_columns=key_columns,
                          reference=csv_reference, add_columns=add_columns)
        with csv_output.open(encoding="utf-8-sig", newline="") as handle:
            cleaned = list(csv.reader(handle))
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Cleaned"
        for row in cleaned:
            sheet.append(row)
        staged = folder / "cleaned.xlsx"
        workbook.save(staged)
        created = []
        try:
            for target, data in ((destination, staged.read_bytes()),
                                 (report, Path(stats["report"]).read_bytes())):
                with target.open("xb") as handle:
                    created.append(target)
                    handle.write(data)
        except Exception:
            for path in created:
                path.unlink(missing_ok=True)
            raise
    stats["report"] = str(report)
    return stats


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("destination")
    parser.add_argument("--report")
    parser.add_argument("--key", nargs="+", metavar="COLUMN")
    parser.add_argument("--reference")
    parser.add_argument("--add", nargs="+", metavar="COLUMN")
    args = parser.parse_args()
    try:
        for label, value in clean_xlsx(args.source, args.destination, args.report,
                                       args.key, args.reference, args.add).items():
            print(f"{label}: {value}")
    except (OSError, ValueError) as error:
        parser.exit(1, f"Cannot clean Excel workbook: {error}\n")
